"""Tables of the paper, gathered from the result files into one Excel workbook.

Reads the CSVs the other scripts write (no training) and writes
results/tables.xlsx with one sheet per table of the paper draft, named by its
number:

  table2_projects     methods per project (dataset_stats.csv)
  table3_scenarios    the balancing scenarios and their train sets (results/main_experiment.csv)
  table4_operators    valid augmentation variants per operator (augmented_pool_summary.csv)
  table5_model_size   the ten model sizes, mean over seeds 42-44 (results/model_size.csv)
  table6_scores       test scores per scenario, mean and sd over seeds 0-9 (results/main_experiment.csv)
  table7_cost         training and inference time per scenario (results/main_experiment.csv)

Means and standard deviations sit in separate numeric columns. Precision,
recall and F1 are those of the smelly class.

Usage:
    pip install openpyxl
    python scripts/6_tables.py
"""
import argparse
import csv
import os
import statistics

SCENARIOS = [("none", "27.5", "4", "No balancing"),
             ("none", "27.5", "12", "No balancing, 12 layers"),
             ("head", "27.5", "4", "Head retrained (SMOTE control)"),
             ("weight", "37.7", "4", "Class weighting"),
             ("weight", "50.0", "4", "Class weighting"),
             ("augment", "37.7", "4", "Augmentation"),
             ("augment", "50.0", "4", "Augmentation"),
             ("smote", "37.7", "4", "SMOTE"),
             ("smote", "50.0", "4", "SMOTE")]
METRICS = [("accuracy", "Accuracy"), ("precision", "Precision"), ("recall", "Recall"), ("f1", "F1"),
           ("mcc", "MCC")]


def read(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def mean_sd(values):
    values = [float(v) for v in values]
    return round(statistics.mean(values), 3), round(statistics.stdev(values), 3) if len(values) > 1 else 0.0


def runs(rows, scenario, share, layers):
    return [r for r in rows if r["scenario"] == scenario and r["smell_pct"] == share and r["layers"] == layers]


def table2(stats):
    header = ["Project", "Methods", "Clean", "Smelly", "Smelly (%)", "Median tokens (smelly)",
              "Max tokens (smelly)", "Median tokens (clean)", "Smelly with switch", "Smelly with else-if chain"]
    body = [[r["project"], int(r["total"]), int(r["clean"]), int(r["smell"]), float(r["smell_pct"]),
             float(r["smell_median_tok"]), int(r["smell_max_tok"]), float(r["clean_median_tok"]),
             int(r["smell_with_switch"]), int(r["smell_with_else_if"])] for r in stats]
    return header, body


def table3(results):
    header = ["Scenario", "Smelly share of train (%)", "Encoder layers", "Training", "Train clean", "Train smelly",
              "Weight clean", "Weight smelly"]
    body = []
    for scenario, share, layers, name in SCENARIOS:
        r = runs(results, scenario, share, layers)[0]
        training = "head only (frozen encoder)" if r["stage"] == "head" else "end to end"
        body.append([name, float(share), int(layers), training, int(r["train_clean"]), int(r["train_smell"]),
                     float(r["weight_clean"]), float(r["weight_smell"])])
    return header, body


def table4(summary):
    header = ["Operator", "Smelly valid", "Smelly tried", "Smelly valid (%)", "Clean valid", "Clean tried",
              "Clean valid (%)", "Unchanged", "Syntax error", "Over token budget", "Tool error"]
    by_op = {}
    for r in summary:
        if r["subset"] in ("experiment_smell", "experiment_clean"):
            # The Total row has no operator number; it goes last
            key = (int(r["op_id"]) if r["op_id"] else 99, r["operator"])
            by_op.setdefault(key, {})[r["subset"]] = r
    body = []
    for (_, operator), parts in sorted(by_op.items()):
        s, c = parts["experiment_smell"], parts["experiment_clean"]
        total = lambda k: int(s[k]) + int(c[k])
        body.append([operator, int(s["ok"]), int(s["total"]), round(100 * int(s["ok"]) / int(s["total"]), 1),
                     int(c["ok"]), int(c["total"]), round(100 * int(c["ok"]) / int(c["total"]), 1),
                     total("unchanged"), total("syntax"), total("over_budget"), total("error")])
    return header, body


def table5(model_size):
    header = ["Encoder layers", "Parameters (M)"] + [x for _, name in METRICS for x in (name, f"{name} sd")] + [
        "Training time (s)", "Energy (Wh)", "Emissions (g CO2e)", "Peak GPU memory (GB)"]
    body = []
    for layers in sorted({int(r["layers"]) for r in model_size}, reverse=True):
        rs = [r for r in model_size if int(r["layers"]) == layers]
        row = [layers, round(int(rs[0]["params_total"]) / 1e6, 1)]
        for metric, _ in METRICS:
            row += mean_sd(r[metric] for r in rs)
        energy = [r["energy_kwh"] for r in rs if r["energy_kwh"]]
        row += [round(statistics.mean(float(r["train_seconds"]) for r in rs), 1),
                round(1000 * statistics.mean(float(v) for v in energy), 2) if energy else None,
                round(1000 * statistics.mean(float(r["emissions_kg"]) for r in rs), 2) if energy else None,
                round(statistics.mean(float(r["peak_gpu_mb"]) for r in rs) / 1024, 1)]
        body.append(row)
    return header, body


def table6(results):
    header = ["Scenario", "Smelly share of train (%)", "Encoder layers"] + [
        x for _, name in METRICS for x in (name, f"{name} sd")] + ["Seeds"]
    body = []
    for scenario, share, layers, name in SCENARIOS:
        rs = runs(results, scenario, share, layers)
        row = [name, float(share), int(layers)]
        for metric, _ in METRICS:
            row += mean_sd(r[metric] for r in rs)
        body.append(row + [len(rs)])
    return header, body


def table7(results):
    """Training time per run: SMOTE and the head control reuse the encoder of the no-balancing run of the
    same seed, so their total adds that run's time to their head stage."""
    header = ["Scenario", "Smelly share of train (%)", "Encoder layers", "Own training (s)", "Own training sd",
              "Total training incl. reused encoder (s)", "Inference on 70 methods (s)", "Peak GPU memory (GB)"]
    base = {(r["seed"], r["layers"]): float(r["train_seconds"]) for r in results if r["scenario"] == "none"}
    body = []
    for scenario, share, layers, name in SCENARIOS:
        rs = runs(results, scenario, share, layers)
        own = mean_sd(r["train_seconds"] for r in rs)
        total = statistics.mean(float(r["train_seconds"]) + (base[(r["seed"], r["layers"])] if r["stage"] == "head"
                                                             else 0.0) for r in rs)
        gpu = [float(r["peak_gpu_mb"]) for r in rs if r["peak_gpu_mb"]]
        body.append([name, float(share), int(layers), own[0], own[1], round(total, 1),
                     round(statistics.mean(float(r["test_seconds"]) for r in rs), 2),
                     round(statistics.mean(gpu) / 1024, 2) if gpu else None])
    return header, body


def write(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, body) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(header)
        for row in body:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, column in enumerate(ws.columns, start=1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--dataset-stats", default="dataset_stats.csv")
    parser.add_argument("--pool-summary", default="augmented_pool_summary.csv")
    parser.add_argument("--model-size", default="results/model_size.csv")
    parser.add_argument("--results", default="results/main_experiment.csv")
    parser.add_argument("--representation", choices=["comment", "nocomment"], default="comment")
    parser.add_argument("--out", default="results/tables.xlsx")
    args = parser.parse_args()

    results = [r for r in read(args.results) if r["representation"] == args.representation]
    model_size = [r for r in read(args.model_size) if r["representation"] == args.representation]
    sheets = {"table2_projects": table2(read(args.dataset_stats)),
              "table3_scenarios": table3(results),
              "table4_operators": table4(read(args.pool_summary)),
              "table5_model_size": table5(model_size),
              "table6_scores": table6(results),
              "table7_cost": table7(results)}
    write(args.out, sheets)
    print(f"-> {args.out}: {', '.join(sheets)}")


if __name__ == "__main__":
    main()
